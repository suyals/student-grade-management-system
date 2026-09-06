from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSETS_DIR = ROOT / "report_assets"
CLEAN_DIR = ASSETS_DIR / "clean"
OUTPUT_DOCX = ROOT / "Shraddha_Suyal_Student_Grade_Management_System_Report.docx"

THEME_BLUE = RGBColor(25, 56, 102)
THEME_MUTED = RGBColor(90, 90, 90)
THEME_LIGHT = RGBColor(228, 235, 247)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def save_cropped_image(source_name: str, target_name: str, box: tuple[int, int, int, int]) -> Path:
    source = Image.open(ASSETS_DIR / source_name).convert("RGB")
    cropped = source.crop(box)
    bordered = ImageOps.expand(cropped, border=2, fill=(220, 228, 240))
    target = CLEAN_DIR / target_name
    bordered.save(target, quality=95)
    return target


def build_architecture_diagram() -> Path:
    width, height = 1600, 950
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(44, bold=True)
    box_title_font = load_font(28, bold=True)
    body_font = load_font(22)
    small_font = load_font(20)

    def rounded_box(x1, y1, x2, y2, title, lines, fill):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=26, fill=fill, outline="#244b8a", width=4)
        draw.text((x1 + 24, y1 + 18), title, font=box_title_font, fill="#173766")
        text_y = y1 + 70
        for line in lines:
            draw.text((x1 + 24, text_y), line, font=small_font, fill="#2f3b4b")
            text_y += 34

    def arrow(start, end):
        draw.line((start, end), fill="#244b8a", width=8)
        ex, ey = end
        draw.polygon(
            [
                (ex, ey),
                (ex - 18, ey - 12),
                (ex - 18, ey + 12),
            ],
            fill="#244b8a",
        )

    draw.text((50, 40), "Architecture Diagram - Student Grade Management System", font=title_font, fill="#173766")
    draw.text(
        (52, 102),
        "A lightweight client-side grading workflow with a static Node.js server and browser-based data persistence.",
        font=body_font,
        fill="#566273",
    )

    rounded_box(90, 260, 360, 470, "User", ["Adds student data", "Enters marks", "Reviews grades"], "#eef4ff")
    rounded_box(
        470,
        220,
        1020,
        520,
        "Browser Interface",
        ["index.html defines the layout", "styles.css provides the responsive dashboard UI", "script.js handles forms, metrics, and table updates"],
        "#f7fbff",
    )
    rounded_box(
        1100,
        220,
        1490,
        430,
        "Grade Logic",
        ["Average calculation", "Letter-grade mapping", "Pass/fail status updates"],
        "#eef8f1",
    )
    rounded_box(
        1100,
        500,
        1490,
        720,
        "Local Storage",
        ["Stores student records", "Retains grades between sessions", "Feeds dashboard reload state"],
        "#fff7eb",
    )
    rounded_box(
        470,
        610,
        1020,
        820,
        "Node.js Static Server",
        ["server.js serves HTML, CSS, and JavaScript", "Provides a simple localhost environment for testing"],
        "#f4efff",
    )

    arrow((360, 365), (470, 365))
    arrow((1020, 340), (1100, 340))
    arrow((1295, 430), (1295, 500))
    arrow((1020, 700), (1100, 700))
    arrow((740, 610), (740, 520))

    draw.text((560, 860), "Figure 1. High-level flow of the Student Grade Management System.", font=body_font, fill="#3e4b5c")

    target = CLEAN_DIR / "architecture-diagram.png"
    image.save(target, quality=95)
    return target


def prepare_assets() -> dict[str, Path]:
    CLEAN_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "home_overview": save_cropped_image("01-empty-dashboard.png", "home-overview.png", (40, 20, 980, 560)),
        "student_registration": save_cropped_image("02-add-student-form.png", "student-registration.png", (40, 160, 980, 520)),
        "marks_assignment": save_cropped_image("03-assign-marks.png", "marks-assignment.png", (40, 500, 980, 920)),
        "dashboard_summary": save_cropped_image("04-performance-report.png", "dashboard-summary.png", (40, 820, 980, 1040)),
        "performance_table": save_cropped_image("04-performance-report.png", "performance-table.png", (40, 1040, 980, 1445)),
    }
    assets["architecture_diagram"] = build_architecture_diagram()
    return assets


def set_run_font(run, name: str, size: int, bold: bool = False, color: RGBColor | None = None, italic: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=0, line_spacing=1.15):
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def add_title_paragraph(doc: Document, text: str, size: int = 16):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8, space_before=0, line_spacing=1.0)
    run = paragraph.add_run(text)
    set_run_font(run, "Times New Roman", size, bold=True, color=THEME_BLUE)
    return paragraph


def add_body_paragraph(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8, space_before=0, line_spacing=1.15)
    run = paragraph.add_run(text)
    set_run_font(run, "Times New Roman", 12)
    return paragraph


def add_bullet_list(doc: Document, items: list[str]):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0, line_spacing=1.1)
        run = paragraph.add_run(item)
        set_run_font(run, "Times New Roman", 12)


def add_numbered_list(doc: Document, items: list[str]):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0, line_spacing=1.1)
        run = paragraph.add_run(item)
        set_run_font(run, "Times New Roman", 12)


def set_cell_background(cell, color: str):
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell_properties.append(shading)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "D7DEE9")
        borders.append(border)
    tbl_pr.append(borders)


def add_standard_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    set_table_borders(table)

    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        cell = header_cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_background(cell, "EAF0FA")
        paragraph = cell.paragraphs[0]
        style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, line_spacing=1.0)
        run = paragraph.add_run(header)
        set_run_font(run, "Times New Roman", 11, bold=True, color=THEME_BLUE)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.LEFT if idx != len(row) - 1 else WD_ALIGN_PARAGRAPH.CENTER, space_after=2, line_spacing=1.0)
            run = paragraph.add_run(value)
            set_run_font(run, "Times New Roman", 10)
    return table


def add_code_block(doc: Document, title: str, code: str):
    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4, space_before=4, line_spacing=1.0)
    run = subtitle.add_run(title)
    set_run_font(run, "Times New Roman", 12, bold=True, color=THEME_BLUE)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    set_table_borders(table)
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.3)
    set_cell_background(cell, "F7F9FC")
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0, space_before=0, line_spacing=1.0)
    run = paragraph.add_run(code)
    set_run_font(run, "Consolas", 9)


def add_image_with_caption(doc: Document, image_path: Path, width: float, caption: str):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, space_before=4, line_spacing=1.0)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width))

    caption_paragraph = doc.add_paragraph()
    style_paragraph(caption_paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, space_before=0, line_spacing=1.0)
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, "Times New Roman", 10, italic=True, color=THEME_MUTED)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_text)
    run._r.append(fld_separate)
    run._r.append(fld_text)
    run._r.append(fld_end)


def set_document_defaults(doc: Document):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25.4)
    section.bottom_margin = Mm(25.4)
    section.left_margin = Mm(25.4)
    section.right_margin = Mm(25.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)

    footer_paragraph = section.footer.paragraphs[0]
    style_paragraph(footer_paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0, line_spacing=1.0)
    run = footer_paragraph.add_run("Page ")
    set_run_font(run, "Times New Roman", 10, color=THEME_MUTED)
    add_page_field(footer_paragraph)


def add_cover_page(doc: Document):
    for _ in range(3):
        doc.add_paragraph("")

    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10, line_spacing=1.0)
    run = paragraph.add_run("Student Grade Management System")
    set_run_font(run, "Times New Roman", 22, bold=True, color=THEME_BLUE)

    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18, line_spacing=1.0)
    run = paragraph.add_run("A PROJECT REPORT")
    set_run_font(run, "Times New Roman", 16, bold=True)

    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8, line_spacing=1.0)
    run = paragraph.add_run("Submitted by")
    set_run_font(run, "Times New Roman", 13)

    table = add_standard_table(
        doc,
        ["Sr No", "Name", "Registration No"],
        [["1", "Shraddha Suyal", ""]],
        [0.9, 3.5, 2.0],
    )
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for _ in range(2):
        doc.add_paragraph("")

    cover_lines = [
        "in partial fulfillment for the award of the degree",
        "of",
        "BACHELOR OF TECHNOLOGY",
        "School of Computer Science and Engineering",
        "August 2026",
    ]
    for idx, line in enumerate(cover_lines):
        paragraph = doc.add_paragraph()
        style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4 if idx < 3 else 6, line_spacing=1.0)
        run = paragraph.add_run(line)
        set_run_font(run, "Times New Roman", 12 if idx not in (2,) else 14, bold=idx == 2)

    doc.add_page_break()


def add_toc_page(doc: Document):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, line_spacing=1.0)
    run = paragraph.add_run("TABLE OF CONTENTS")
    set_run_font(run, "Times New Roman", 16, bold=True, color=THEME_BLUE)

    toc_rows = [
        ["1", "Abstract", "3"],
        ["2", "Introduction", "4"],
        ["3", "Problem Statement", "5"],
        ["4", "Objectives", "6"],
        ["5", "Scope", "7"],
        ["6", "System Design & Analysis", "8"],
        ["7", "Technologies Used", "9"],
        ["8", "Architect Diagram", "10"],
        ["9", "Module Description", "11"],
        ["10", "Implementation", "11"],
        ["11", "Code Snippets", "12-13"],
        ["12", "Screenshot of Web Pages", "14-16"],
        ["13", "Testing", "17"],
        ["14", "Result", "18"],
        ["15", "Output Screenshot", "18"],
        ["16", "Conclusion", "19"],
        ["17", "Future Work", "19"],
        ["18", "References", "20"],
    ]
    add_standard_table(doc, ["Sr No", "Topic", "Page"], toc_rows, [1.0, 4.6, 1.2])
    doc.add_page_break()


def build_report():
    assets = prepare_assets()
    doc = Document()
    set_document_defaults(doc)

    add_cover_page(doc)
    add_toc_page(doc)

    add_title_paragraph(doc, "ABSTRACT")
    add_body_paragraph(
        doc,
        "The Student Grade Management System is a web-based academic dashboard developed to simplify the process of recording student details, entering subject marks, and generating instant grade-based performance summaries. The system allows a teacher or coordinator to register students with roll numbers, assign marks for Mathematics, Science, and English, and automatically compute each student's average score, grade, and pass or fail status. It also highlights class-level analytics such as total students, topper details, and class average so that the overall performance picture is available at a glance.",
    )
    add_body_paragraph(
        doc,
        "The project is built using HTML5, CSS3, JavaScript, and a lightweight Node.js static server for local execution. Browser local storage is used to preserve records between sessions, while the JavaScript application logic handles validation, grade mapping, and dynamic dashboard rendering in real time. The result is a simple, responsive, and practical classroom utility that demonstrates how front-end web technologies can be used to digitize routine academic record management tasks.",
    )
    doc.add_page_break()

    add_title_paragraph(doc, "INTRODUCTION")
    add_body_paragraph(
        doc,
        "Managing student performance manually through paper registers or spreadsheets can be time-consuming, repetitive, and error-prone. Teachers often need to maintain lists of students, enter marks for multiple subjects, calculate averages, assign grades, identify toppers, and review class-level performance. When these tasks are handled manually, even small arithmetic mistakes or missing entries can affect the accuracy of the final report.",
    )
    add_body_paragraph(
        doc,
        "The Student Grade Management System addresses this challenge by offering a focused academic dashboard that combines student registration, marks entry, and performance reporting in a single interface. The application is designed for quick everyday use, with clear forms, automatic calculations, and a live data table that updates instantly whenever a new student is added or marks are modified.",
    )
    add_body_paragraph(
        doc,
        "From a technical perspective, the project demonstrates a clean front-end architecture in which HTML defines the structure, CSS provides the dashboard styling, and JavaScript coordinates all validation, calculations, state updates, and persistence. A small Node.js server is used to serve the static files locally during execution and testing.",
    )
    doc.add_page_break()

    add_title_paragraph(doc, "PROBLEM STATEMENT")
    add_body_paragraph(
        doc,
        "In many academic settings, student records are still maintained using disconnected tools such as notebooks, spreadsheets, or multiple paper files. This makes the process of registering students, recording subject-wise marks, and producing quick performance summaries inefficient. Repeated manual calculations of average marks and grade categories also increase the chance of mistakes and consume time that could otherwise be used for teaching or student support.",
    )
    add_body_paragraph(
        doc,
        "The Student Grade Management System is designed to solve this problem by providing a single digital interface for entering student information and marks while automatically generating averages, letter grades, pass or fail status, topper information, and class-level statistics. The project reduces repetitive work, improves data clarity, and demonstrates a structured approach to small-scale academic record management through a browser-based dashboard.",
    )
    doc.add_page_break()

    add_title_paragraph(doc, "OBJECTIVES")
    add_bullet_list(
        doc,
        [
            "To register students with essential academic identifiers such as student name and roll number.",
            "To provide a subject-wise marks entry interface for Mathematics, Science, and English.",
            "To calculate average marks automatically and map them to grade categories without manual intervention.",
            "To display pass or fail status immediately after marks are updated.",
            "To present summary indicators such as total students, topper, and class average in a dashboard format.",
            "To preserve student records in the browser so that the application remains useful across sessions.",
        ],
    )
    doc.add_page_break()

    add_title_paragraph(doc, "SCOPE")
    add_body_paragraph(
        doc,
        "The scope of the Student Grade Management System is centered on small to medium classroom use where a teacher, mentor, or coordinator needs a quick and reliable method to maintain student performance data. The current version supports student registration, subject-wise marks entry for three subjects, automatic average calculation, grade assignment, pass or fail determination, and live performance reporting through a responsive dashboard.",
    )
    add_body_paragraph(
        doc,
        "The application is intentionally lightweight and easy to deploy because it does not depend on a database server or external APIs. Its current implementation is best suited for local use, classroom demonstrations, and introductory academic management scenarios. Future enhancements can extend the project toward multi-subject support, exportable reports, authentication, cloud-backed storage, attendance integration, and richer analytics such as subject trends and student comparisons.",
    )
    doc.add_page_break()

    add_title_paragraph(doc, "SYSTEM DESIGN & ANALYSIS")
    add_body_paragraph(
        doc,
        "The Student Grade Management System follows a simple client-side architecture in which the browser hosts the full user interface and application logic, while a minimal Node.js server delivers the static project files. The application interface is divided into four major visual blocks: a hero header, a student registration panel, a marks entry panel, and a reporting area containing summary cards and the performance table.",
    )
    add_body_paragraph(
        doc,
        "The system workflow begins when the user enters student details through the registration form. After validation checks for empty fields and duplicate roll numbers, the student record is added to the in-memory student array and stored in browser local storage. Once a student is available in the list, the marks module allows subject scores to be assigned. JavaScript then calculates the average score, determines the grade through conditional rules, and updates the pass or fail status before the interface is re-rendered.",
    )
    add_body_paragraph(
        doc,
        "The reporting area continuously reflects the latest state of the data. It shows the total number of students, identifies the topper by sorting student averages, calculates the class average, and lists all student records in a tabular view. This design keeps the application responsive and easy to understand, while also making the codebase suitable for learning core front-end development patterns such as form handling, DOM updates, data validation, and local persistence.",
    )
    doc.add_page_break()

    add_title_paragraph(doc, "TECHNOLOGIES USED")
    add_body_paragraph(
        doc,
        "The project uses a compact set of front-end and Node.js technologies. Each technology contributes a specific responsibility to the interface, interaction logic, or local execution environment.",
    )
    add_standard_table(
        doc,
        ["Technology", "Purpose"],
        [
            ["HTML5", "Provides the semantic structure for forms, summary cards, and the performance table."],
            ["CSS3", "Creates the dashboard layout, gradients, cards, responsive grid, and visual styling."],
            ["JavaScript (ES6)", "Handles validation, state updates, grade calculation, rendering, and user feedback messages."],
            ["Browser Local Storage", "Preserves student records and marks between page reloads on the same device."],
            ["Node.js", "Runs the local static server used to preview and test the application in the browser."],
            ["HTTP, FS, and Path Modules", "Serve project files safely and map requests to the correct assets in server.js."],
            ["DOM APIs and crypto.randomUUID()", "Support dynamic table rendering, form handling, and stable student identifiers."],
        ],
        [1.8, 4.7],
    )
    doc.add_page_break()

    add_title_paragraph(doc, "ARCHITECT DIAGRAM")
    add_body_paragraph(
        doc,
        "Figure 1 illustrates the high-level architecture of the Student Grade Management System and the relationship between the user, browser interface, grading logic, persistent browser storage, and the Node.js static server.",
    )
    add_image_with_caption(doc, assets["architecture_diagram"], 6.2, "Figure 1. Architecture diagram of the Student Grade Management System.")
    doc.add_page_break()

    add_title_paragraph(doc, "MODULE DESCRIPTION")
    add_numbered_list(
        doc,
        [
            "Student Registration Module: captures the student name and roll number, validates required fields, and prevents duplicate roll entries.",
            "Marks Entry Module: allows the user to select a registered student and enter marks for Mathematics, Science, and English.",
            "Grade Evaluation Module: calculates the average score, assigns the grade band, and updates the pass or fail status.",
            "Summary Dashboard Module: computes the total number of students, topper information, and class average.",
            "Performance Report Module: renders a table view of every student record with subject scores and final outcome details.",
        ],
    )
    add_title_paragraph(doc, "IMPLEMENTATION", size=14)
    add_body_paragraph(
        doc,
        "The implementation is organized into four project files. The `index.html` file structures the dashboard and form controls. The `styles.css` file defines the dark themed layout, responsive grid behavior, and reusable card treatments. The `script.js` file contains the complete application logic for form submission, validation, grade calculation, summary generation, local storage persistence, and dynamic table rendering.",
    )
    add_body_paragraph(
        doc,
        "The `server.js` file provides a lightweight Node.js HTTP server so that the project can be executed consistently through `http://127.0.0.1:8000/`. It resolves requests safely, determines the correct MIME type, and serves the static assets required by the dashboard. This separation keeps the project easy to understand while still reflecting a practical web development workflow.",
    )
    doc.add_page_break()

    add_title_paragraph(doc, "CODE SNIPPETS")
    add_code_block(
        doc,
        "Snippet 1. Grade mapping and metrics update logic",
        """function calculateGrade(average) {
  if (average >= 90) return "A+";
  if (average >= 80) return "A";
  if (average >= 70) return "B";
  if (average >= 60) return "C";
  if (average >= 50) return "D";
  return "F";
}

function updateStudentMetrics(student) {
  const average = calculateAverage(student.marks);
  student.average = Number(average.toFixed(1));
  student.grade = calculateGrade(student.average);
  student.status = student.average >= 40 ? "Pass" : "Fail";
}""",
    )
    add_code_block(
        doc,
        "Snippet 2. Student registration workflow",
        """studentForm.onsubmit = function (event) {
  event.preventDefault();

  const name = document.getElementById("studentName").value.trim();
  const rollNumber = document.getElementById("rollNumber").value.trim();

  if (!name || !rollNumber) {
    showMessage(studentMessage, "Please enter both student name and roll number.", true);
    return;
  }

  const existingRoll = students.some((student) => student.rollNumber === rollNumber);
  if (existingRoll) {
    showMessage(studentMessage, "A student with this roll number already exists.", true);
    return;
  }
};""",
    )
    doc.add_page_break()

    add_title_paragraph(doc, "CODE SNIPPETS")
    add_code_block(
        doc,
        "Snippet 3. Dynamic table rendering",
        """function renderStudents() {
  studentTableBody.innerHTML = "";
  studentSelect.innerHTML = '<option value="">Choose a student</option>';

  students.forEach((student) => {
    const row = document.createElement("tr");
    row.innerHTML = `
      <td>${student.name}</td>
      <td>${student.rollNumber}</td>
      <td>${student.marks.math}</td>
      <td>${student.marks.science}</td>
      <td>${student.marks.english}</td>
      <td>${student.average.toFixed(1)}</td>
      <td>${student.grade}</td>
      <td class="${student.status === "Pass" ? "status-pass" : "status-fail"}">${student.status}</td>
    `;
    studentTableBody.appendChild(row);
  });
}""",
    )
    add_code_block(
        doc,
        "Snippet 4. Static Node.js file server",
        """const server = http.createServer((req, res) => {
  const requestUrl = new URL(req.url, 'http://127.0.0.1');
  const pathname = requestUrl.pathname === '/' ? '/index.html' : requestUrl.pathname;
  const safePath = path.normalize(pathname).replace(/^([a-zA-Z]:[\\\\/])?/, '');
  const filePath = path.join(rootDir, safePath);

  if (!filePath.startsWith(rootDir)) {
    res.writeHead(403, { 'Content-Type': 'text/plain; charset=utf-8' });
    res.end('Forbidden');
    return;
  }
});""",
    )
    doc.add_page_break()

    add_title_paragraph(doc, "SCREENSHOTS OF WEB PAGES")
    add_image_with_caption(doc, assets["home_overview"], 6.0, "1. Home page overview of the academic dashboard.")
    add_image_with_caption(doc, assets["student_registration"], 6.0, "2. Student registration panel with name and roll number inputs.")
    doc.add_page_break()

    add_image_with_caption(doc, assets["marks_assignment"], 6.0, "3. Marks assignment module showing subject-wise entry fields.")
    add_image_with_caption(doc, assets["dashboard_summary"], 6.0, "4. Summary cards highlighting total students, topper, and class average.")
    doc.add_page_break()

    add_image_with_caption(doc, assets["performance_table"], 6.2, "5. Student performance table displaying grades and final status.")
    doc.add_page_break()

    add_title_paragraph(doc, "TESTING")
    add_body_paragraph(
        doc,
        "Functional testing was performed on the primary modules of the Student Grade Management System to confirm that the input flow, calculations, and output panels behave as expected.",
    )
    add_standard_table(
        doc,
        ["Test Case", "Expected Result", "Status"],
        [
            ["Student registration", "A new student record is created with valid name and roll number.", "Passed"],
            ["Duplicate roll validation", "Duplicate roll numbers are rejected with an error message.", "Passed"],
            ["Marks entry", "The selected student's marks are stored successfully.", "Passed"],
            ["Average calculation", "Average score is calculated to one decimal place.", "Passed"],
            ["Grade generation", "The correct grade is assigned according to the defined score bands.", "Passed"],
            ["Pass or fail status", "Student status updates automatically after marks are saved.", "Passed"],
            ["Summary cards", "Topper and class average values refresh after each update.", "Passed"],
            ["Responsive layout", "The dashboard remains usable across smaller screen widths.", "Passed"],
        ],
        [2.2, 3.7, 1.0],
    )
    doc.add_page_break()

    add_title_paragraph(doc, "RESULT")
    add_body_paragraph(
        doc,
        "The Student Grade Management System was successfully developed and tested as a responsive classroom dashboard. The application correctly supports student registration, subject-wise marks entry, average calculation, grade assignment, and live performance reporting. The final interface presents both student-level and class-level insights in a clear format, which improves the speed and reliability of basic academic record management.",
    )
    add_title_paragraph(doc, "OUTPUT SCREENSHOT", size=14)
    add_image_with_caption(doc, assets["dashboard_summary"], 6.0, "Final dashboard summary after completing student and marks entry.")
    doc.add_page_break()

    add_title_paragraph(doc, "CONCLUSION")
    add_body_paragraph(
        doc,
        "The Student Grade Management System demonstrates how a focused web application can simplify an everyday academic task. By combining student registration, marks management, grade computation, summary analytics, and a live reporting table in one interface, the project reduces manual work and improves the clarity of classroom performance tracking.",
    )
    add_body_paragraph(
        doc,
        "The project also reflects core front-end development concepts such as structured HTML, responsive CSS, DOM manipulation, validation, local persistence, and modular JavaScript logic. Its clean architecture makes it suitable both as a practical classroom utility and as a learning project for web development fundamentals.",
    )
    add_title_paragraph(doc, "FUTURE WORK", size=14)
    add_bullet_list(
        doc,
        [
            "Support additional subjects and customizable grading schemes.",
            "Add edit and delete actions for existing student records.",
            "Export the final report to PDF, CSV, or printable mark sheets.",
            "Integrate authentication and role-based access for teachers and administrators.",
            "Connect the application to a database for multi-device access and long-term record storage.",
        ],
    )
    doc.add_page_break()

    add_title_paragraph(doc, "REFERENCES")
    references = [
        "1. MDN Web Docs. HTML basics. Available at: https://developer.mozilla.org/en-US/docs/Learn/HTML/Introduction_to_HTML",
        "2. MDN Web Docs. CSS first steps. Available at: https://developer.mozilla.org/en-US/docs/Learn/CSS/First_steps",
        "3. MDN Web Docs. JavaScript guide. Available at: https://developer.mozilla.org/en-US/docs/Web/JavaScript/Guide",
        "4. MDN Web Docs. Web Storage API. Available at: https://developer.mozilla.org/en-US/docs/Web/API/Web_Storage_API",
        "5. MDN Web Docs. Crypto.randomUUID(). Available at: https://developer.mozilla.org/en-US/docs/Web/API/Crypto/randomUUID",
        "6. Node.js Documentation. HTTP module. Available at: https://nodejs.org/api/http.html",
        "7. Node.js Documentation. File system module. Available at: https://nodejs.org/api/fs.html",
        "8. Node.js Documentation. Path module. Available at: https://nodejs.org/api/path.html",
    ]
    for item in references:
        paragraph = doc.add_paragraph()
        style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=5, line_spacing=1.1)
        run = paragraph.add_run(item)
        set_run_font(run, "Times New Roman", 11)

    doc.save(OUTPUT_DOCX)
    print(f"Saved report to: {OUTPUT_DOCX}")


if __name__ == "__main__":
    build_report()
